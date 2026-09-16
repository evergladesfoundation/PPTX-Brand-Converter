#!/usr/bin/env python3
"""Rebuild OUTPUT.pptx from plan.json + template. Never deepcopy charts or paint chrome."""

from __future__ import annotations

import argparse
import io
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from colorways import (  # noqa: E402
    apply_colorway,
    duplicate_slide,
    fill_text_slots,
    patch_theme_part,
    place_source_photos,
    remap_srgb_in_slide,
    strip_pic_locks,
    unlock_slide_pictures,
)
from helpers import (  # noqa: E402
    add_rebuilt_chart,
    apply_line,
    apply_picture_crop,
    apply_round_rect,
    apply_run_style,
    contain_fit,
    content_safe_box,
    crop_cut_fraction,
    delete_all_slides,
    delete_slide,
    delete_shape,
    dump_json,
    ensure_pptx,
    enum_name,
    is_screenshot_blob,
    layout_index_for_role,
    load_json,
    maybe_recompress,
    parse_brand_md,
    recreate_sections,
    reduce_title_size,
    scale_geometry,
    set_slide_hidden,
    strip_effect_list,
    style_table,
    theme_color,
    tint,
    unused_placeholder,
    uses_curly_quotes,
    write_table_cell_text,
    write_text_frame,
    is_totals_row,
)


def find_placeholder(slide, type_names: set[str], idx: int | None = None):
    matches = []
    for shape in slide.placeholders:
        try:
            name = enum_name(shape.placeholder_format.type).upper()
        except Exception:
            continue
        if name in type_names:
            matches.append(shape)
    if idx is not None:
        for shape in matches:
            try:
                if int(shape.placeholder_format.idx) == int(idx):
                    return shape
            except Exception:
                continue
    if not matches:
        return None
    matches.sort(key=lambda s: int(s.left or 0))
    return matches[0]


def fill_title(slide, placement: dict[str, Any], entry: dict[str, Any], tokens: dict[str, Any], brand: dict[str, Any]) -> None:
    ph_info = placement.get("placeholder") or {}
    type_names = {
        enum_name(ph_info.get("type_name") or "TITLE").upper(),
        "TITLE",
        "CENTER_TITLE",
        "VERTICAL_TITLE",
    }
    shape = find_placeholder(slide, type_names, ph_info.get("idx"))
    text = placement.get("text") or ""
    if not shape:
        box = ph_info if ph_info.get("width") else content_safe_box(tokens, tokens.get("_layout_index"))
        shape = slide.shapes.add_textbox(
            int(box.get("left") or 0),
            int(box.get("top") or 0),
            int(box.get("width") or tokens["slide_width"]),
            int(min(box.get("height") or 1, int(tokens["slide_height"] * 0.18) or 1)),
        )
    kind = "cover_title" if entry["role"] == "TITLE" else ("section_title" if entry["role"] == "SECTION" else "title")
    start = int(placement.get("font_pt") or (40 if kind == "cover_title" else 36 if kind == "section_title" else 28))
    box_w = int(shape.width)
    box_h = int(shape.height)
    size = reduce_title_size(text, box_w, box_h, start)
    write_text_frame(
        shape.text_frame,
        [{"text": text, "level": 0, "alignment": "left" if kind != "cover_title" else "center", "runs": [{"text": text}]}],
        tokens,
        brand,
        kind=kind,
        role=entry["role"],
        title_size=size,
        curly_quotes=uses_curly_quotes(tokens),
    )


def fill_paragraph_target(slide, placement: dict[str, Any], tokens: dict[str, Any], brand: dict[str, Any], kind: str, role: str) -> None:
    paras = placement.get("paragraphs")
    if not paras and placement.get("text"):
        paras = [{"text": placement["text"], "level": 0, "runs": [{"text": placement["text"]}]}]
    paras = paras or []
    ph_info = placement.get("placeholder")
    shape = None
    if placement.get("target") == "placeholder" and ph_info:
        type_name = enum_name(ph_info.get("type_name") or "BODY").upper()
        names = {type_name}
        if type_name in {"BODY", "OBJECT"}:
            names.update({"BODY", "OBJECT", "VERTICAL_BODY", "SUBTITLE"})
        if type_name == "SUBTITLE":
            names.add("SUBTITLE")
        shape = find_placeholder(slide, names, ph_info.get("idx"))
        if shape is None:
            # Prefer left-to-right body order: pick first unused body-like placeholder.
            for candidate in sorted(slide.placeholders, key=lambda s: int(s.left or 0)):
                try:
                    t = enum_name(candidate.placeholder_format.type).upper()
                except Exception:
                    continue
                if t in names and unused_placeholder(candidate):
                    shape = candidate
                    break
    if shape is not None and getattr(shape, "has_text_frame", False):
        write_text_frame(
            shape.text_frame,
            paras,
            tokens,
            brand,
            kind=kind,
            role=role,
            body_size=placement.get("font_pt"),
            curly_quotes=uses_curly_quotes(tokens),
        )
        return
    box = placement.get("placeholder") or placement.get("bounds") or content_safe_box(tokens, None)
    if placement.get("bounds"):
        box = scale_geometry(placement["bounds"], {"width": tokens.get("_src_w"), "height": tokens.get("_src_h")}, tokens, tokens.get("_layout_index"))
    tb = slide.shapes.add_textbox(int(box["left"]), int(box.get("top", 0)), int(box.get("width", 1)), int(box.get("height", 1)))
    write_text_frame(
        tb.text_frame,
        paras,
        tokens,
        brand,
        kind=kind,
        role=role,
        body_size=placement.get("font_pt"),
        curly_quotes=uses_curly_quotes(tokens),
    )


def add_picture_placement(slide, placement: dict[str, Any], tokens: dict[str, Any], brand: dict[str, Any], flags: list[str]) -> None:
    image = placement.get("image") or {}
    path = image.get("path")
    if not path or not Path(path).exists():
        return
    blob = Path(path).read_bytes()
    blob, ext = maybe_recompress(blob, image.get("ext") or "png")
    px = image.get("px") or [None, None]
    img_w, img_h = px[0], px[1]
    if not img_w or not img_h:
        try:
            from PIL import Image

            with Image.open(io.BytesIO(blob)) as im:
                img_w, img_h = im.size
        except Exception:
            img_w, img_h = int(placement.get("bounds", {}).get("width") or 1), int(placement.get("bounds", {}).get("height") or 1)
    crop = image.get("crop") or {}
    target_box = None
    if placement.get("target") == "picture_placeholder" and placement.get("placeholder"):
        ph = placement["placeholder"]
        target_box = {"left": int(ph["left"]), "top": int(ph["top"]), "width": int(ph["width"]), "height": int(ph["height"])}
        pic_ph = find_placeholder(slide, {"PICTURE", "MEDIA_CLIP", "OBJECT"}, ph.get("idx"))
        if pic_ph is not None and crop_cut_fraction(crop) <= 0.15:
            try:
                inserted = pic_ph.insert_picture(io.BytesIO(blob))
                apply_picture_crop(inserted, crop)
                strip_effect_list(inserted)
                strip_pic_locks(inserted._element)
                _post_style_picture(inserted, blob, ext, img_w, img_h, tokens, brand)
                return
            except Exception:
                pass
            target_box = {"left": int(pic_ph.left), "top": int(pic_ph.top), "width": int(pic_ph.width), "height": int(pic_ph.height)}
    if target_box is None:
        if placement.get("bounds"):
            target_box = scale_geometry(
                placement["bounds"],
                {"width": tokens.get("_src_w"), "height": tokens.get("_src_h")},
                tokens,
                tokens.get("_layout_index"),
            )
        else:
            target_box = content_safe_box(tokens, tokens.get("_layout_index"))
    fit = contain_fit(int(img_w), int(img_h), target_box)
    picture = slide.shapes.add_picture(io.BytesIO(blob), fit["left"], fit["top"], fit["width"], fit["height"])
    strip_pic_locks(picture._element)
    apply_picture_crop(picture, crop)
    strip_effect_list(picture)
    _post_style_picture(picture, blob, ext, img_w, img_h, tokens, brand)


def _post_style_picture(picture, blob: bytes, ext: str, img_w: int, img_h: int, tokens: dict[str, Any], brand: dict[str, Any]) -> None:
    if brand.get("rounded_corners"):
        apply_round_rect(picture)
    if brand.get("image_border"):
        apply_line(picture, theme_color(tokens, "accent1"), 0.75)
    aspect = (img_w / max(1, img_h)) if img_w and img_h else 1.0
    if is_screenshot_blob(blob, ext, aspect):
        apply_line(picture, tint(theme_color(tokens, "dk2"), 0.50), 0.75)


def add_table_placement(slide, placement: dict[str, Any], tokens: dict[str, Any], brand: dict[str, Any]) -> None:
    chunk = placement.get("table_chunk")
    table_info = placement.get("table") or {}
    if chunk:
        header = chunk.get("header") or []
        rows = chunk.get("rows") or []
        cells = [header] + rows if header else rows
        font_pt = int(chunk.get("font_pt") or 11)
    else:
        cells = table_info.get("cells") or []
        font_pt = 11
    if not cells:
        return
    nrows = len(cells)
    ncols = max(len(r) for r in cells)
    safe = content_safe_box(tokens, tokens.get("_layout_index"))
    src_box = placement.get("bounds") or safe
    box = scale_geometry(src_box, {"width": tokens.get("_src_w"), "height": tokens.get("_src_h")}, tokens, tokens.get("_layout_index"))
    # Prefer the full content-safe width for tables.
    box["left"] = safe["left"]
    box["width"] = safe["width"]
    box["top"] = max(safe["top"], box["top"])
    min_h = int(0.28 * 914400) * nrows
    box["height"] = max(min_h, min(box["height"], safe["top"] + safe["height"] - box["top"]))
    graphic = slide.shapes.add_table(nrows, ncols, int(box["left"]), int(box["top"]), int(box["width"]), int(box["height"]))
    table = graphic.table
    widths = table_info.get("col_widths") or []
    if widths and sum(widths) > 0:
        total = sum(widths)
        for i, w in enumerate(widths[:ncols]):
            table.columns[i].width = int(box["width"] * (w / total))
    totals = is_totals_row(cells[-1]) if cells else False
    style_table(table, tokens, brand, totals_row=totals)
    curly = uses_curly_quotes(tokens)
    for r, row in enumerate(cells):
        for c in range(ncols):
            info = row[c] if c < len(row) else {"text": ""}
            write_table_cell_text(
                table.cell(r, c),
                info,
                tokens,
                brand,
                header=r == 0,
                totals=totals and r == nrows - 1,
                font_pt=font_pt if r else 12,
                curly_quotes=curly,
            )
    for merge in table_info.get("merges") or []:
        r, c = int(merge["row"]), int(merge["col"])
        rs, gs = int(merge.get("row_span") or 1), int(merge.get("grid_span") or 1)
        if r == 0 and chunk and chunk.get("cont"):
            continue
        if rs <= 1 and gs <= 1:
            continue
        try:
            end_r = min(nrows - 1, r + rs - 1)
            end_c = min(ncols - 1, c + gs - 1)
            if end_r != r or end_c != c:
                table.cell(r, c).merge(table.cell(end_r, end_c))
        except Exception:
            pass


def add_chart_placement(slide, placement: dict[str, Any], tokens: dict[str, Any], brand: dict[str, Any], flags: list[str], work: dict[str, Any]) -> None:
    chart = placement.get("chart") or {}
    safe = content_safe_box(tokens, tokens.get("_layout_index"))
    box = placement.get("bounds")
    if box:
        box = scale_geometry(box, {"width": tokens.get("_src_w"), "height": tokens.get("_src_h")}, tokens, tokens.get("_layout_index"))
    else:
        box = dict(safe)
    box["left"] = max(box["left"], safe["left"])
    box["top"] = max(box["top"], safe["top"])
    box["width"] = min(box["width"], safe["left"] + safe["width"] - box["left"])
    box["height"] = min(box["height"], safe["top"] + safe["height"] - box["top"])
    if chart.get("external") or not chart.get("supported"):
        flags.append("NEEDS_MANUAL_REBUILD")
        raster = work.get("chart_rasters", {}).get(str(placement.get("shape_id")))
        if raster and Path(raster).exists():
            slide.shapes.add_picture(raster, int(box["left"]), int(box["top"]), int(box["width"]), int(box["height"]))
        else:
            rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(box["left"]), int(box["top"]), int(box["width"]), int(box["height"]))
            rect.fill.solid()
            rect.fill.fore_color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(0xF5, 0xF5, 0xF5)
            apply_line(rect, theme_color(tokens, "dk2"), 0.75)
            tf = rect.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = "Chart needs manual rebuild"
            apply_run_style(
                run,
                tokens,
                brand,
                font_kind="minor",
                size_pt=12,
                color_hex=theme_color(tokens, "dk2"),
                bold=False,
                italic=True,
            )
        return
    add_rebuilt_chart(slide, chart, box, tokens, brand)


def add_callout(slide, placement: dict[str, Any], tokens: dict[str, Any], brand: dict[str, Any]) -> None:
    box = scale_geometry(
        placement.get("bounds") or content_safe_box(tokens, tokens.get("_layout_index")),
        {"width": tokens.get("_src_w"), "height": tokens.get("_src_h")},
        tokens,
        tokens.get("_layout_index"),
    )
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(box["left"]), int(box["top"]), int(box["width"]), int(box["height"]))
    from pptx.dml.color import RGBColor
    from helpers import rgb

    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(tint(theme_color(tokens, "accent1"), 0.88))
    apply_line(shape, theme_color(tokens, "accent1"), 0.75)
    write_text_frame(
        shape.text_frame,
        placement.get("paragraphs") or [],
        tokens,
        brand,
        kind="callout",
        curly_quotes=uses_curly_quotes(tokens),
    )


def add_line(slide, placement: dict[str, Any], tokens: dict[str, Any]) -> None:
    box = scale_geometry(
        placement.get("bounds") or {"left": 0, "top": 0, "width": 1, "height": 1},
        {"width": tokens.get("_src_w"), "height": tokens.get("_src_h")},
        tokens,
        tokens.get("_layout_index"),
    )
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        int(box["left"]),
        int(box["top"]),
        int(box["left"] + box["width"]),
        int(box["top"] + box["height"]),
    )
    apply_line(connector, theme_color(tokens, "dk2"), 1.0)


def add_media_placeholder(slide, placement: dict[str, Any], tokens: dict[str, Any], brand: dict[str, Any], flags: list[str]) -> None:
    flags.append("MEDIA_MISSING")
    box = scale_geometry(
        placement.get("bounds") or content_safe_box(tokens, tokens.get("_layout_index")),
        {"width": tokens.get("_src_w"), "height": tokens.get("_src_h")},
        tokens,
        tokens.get("_layout_index"),
    )
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(box["left"]), int(box["top"]), int(box["width"]), int(box["height"]))
    from helpers import rgb

    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(tint(theme_color(tokens, "dk2"), 0.85))
    apply_line(shape, theme_color(tokens, "dk2"), 0.75)
    tf = shape.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = placement.get("name") or "Media missing"
    apply_run_style(
        run,
        tokens,
        brand,
        font_kind="minor",
        size_pt=12,
        color_hex=theme_color(tokens, "dk1"),
        bold=False,
        italic=True,
    )


def cleanup_placeholders(slide) -> None:
    for shape in list(slide.placeholders):
        if unused_placeholder(shape):
            delete_shape(shape)


def _placement_text(placement: dict[str, Any]) -> str:
    if placement.get("text"):
        return str(placement["text"])
    paras = placement.get("paragraphs") or []
    return "\n".join((p.get("text") or "").strip() for p in paras if (p.get("text") or "").strip())


def fill_prototype_slide(slide, entry: dict[str, Any], tokens: dict[str, Any], brand: dict[str, Any], flags: list[str], work: dict[str, Any]) -> None:
    title = ""
    subtitle = ""
    body_paras: list[str] = []
    deferred: list[dict[str, Any]] = []
    for placement in entry.get("placements") or []:
        kind = placement.get("kind")
        if placement.get("flag"):
            flags.append(placement["flag"])
        if kind == "title":
            title = _placement_text(placement) or title
        elif kind == "subtitle":
            subtitle = _placement_text(placement) or subtitle
        elif kind in {"body", "caption", "footnote", "textbox", "smartart"}:
            if kind == "smartart":
                flags.append("SMARTART_FLATTENED")
            paras = placement.get("paragraphs") or []
            if paras:
                body_paras.extend((p.get("text") or "").strip() for p in paras if (p.get("text") or "").strip())
            elif placement.get("text"):
                body_paras.append(str(placement["text"]))
        else:
            deferred.append(placement)
    fill_text_slots(slide, title, subtitle, body_paras)
    pictures: list[dict[str, Any]] = []
    rest_deferred: list[dict[str, Any]] = []
    for placement in deferred:
        if placement.get("kind") == "picture":
            pictures.append(placement)
        else:
            rest_deferred.append(placement)
    place_source_photos(slide, pictures, tokens, flags)
    for placement in rest_deferred:
        kind = placement.get("kind")
        if kind == "table":
            add_table_placement(slide, placement, tokens, brand)
        elif kind == "chart":
            add_chart_placement(slide, placement, tokens, brand, flags, work)
        elif kind == "callout":
            add_callout(slide, placement, tokens, brand)
        elif kind == "line":
            add_line(slide, placement, tokens)
        elif kind == "media":
            add_media_placeholder(slide, placement, tokens, brand, flags)
    unlock_slide_pictures(slide)


def fill_layout_slide(slide, entry: dict[str, Any], tokens: dict[str, Any], brand: dict[str, Any], flags: list[str], work: dict[str, Any]) -> None:
    for placement in entry.get("placements") or []:
        kind = placement.get("kind")
        if placement.get("flag"):
            flags.append(placement["flag"])
        if kind == "title":
            fill_title(slide, placement, entry, tokens, brand)
        elif kind in {"subtitle", "body"}:
            fill_paragraph_target(slide, placement, tokens, brand, "subtitle" if kind == "subtitle" else "body", entry["role"])
        elif kind in {"caption", "footnote", "textbox"}:
            fill_paragraph_target(slide, placement, tokens, brand, "caption" if kind == "caption" else "textbox", entry["role"])
        elif kind == "picture":
            add_picture_placement(slide, placement, tokens, brand, flags)
        elif kind == "table":
            add_table_placement(slide, placement, tokens, brand)
        elif kind == "chart":
            add_chart_placement(slide, placement, tokens, brand, flags, work)
        elif kind == "smartart":
            flags.append("SMARTART_FLATTENED")
            fill_paragraph_target(slide, placement, tokens, brand, "body", entry["role"])
        elif kind == "callout":
            add_callout(slide, placement, tokens, brand)
        elif kind == "line":
            add_line(slide, placement, tokens)
        elif kind == "media":
            add_media_placeholder(slide, placement, tokens, brand, flags)
    unlock_slide_pictures(slide)


def build_presentation(
    *,
    template_path: Path,
    source_path: Path,
    out_dir: Path,
    tokens: dict[str, Any],
    manifest: dict[str, Any],
    plan: dict[str, Any],
    brand_md: Path | None,
    colorway: str | None = None,
) -> dict[str, Any]:
    tokens = apply_colorway(dict(tokens), colorway or tokens.get("colorway") or plan.get("colorway"))
    file_brand = parse_brand_md(brand_md)
    token_brand = dict(tokens.get("brand_md") or {})
    brand = {**file_brand, **token_brand}
    for key in ("fonts", "color_whitelist", "chart_palette"):
        brand[key] = sorted(set((file_brand.get(key) or []) + (token_brand.get(key) or [])))
    pptx_template = ensure_pptx(template_path, out_dir / "TEMPLATE.pptx") if template_path.suffix.lower() == ".potx" else template_path
    dest = Presentation(str(pptx_template))
    tokens["_src_w"] = int(manifest.get("slide_width") or tokens["slide_width"])
    tokens["_src_h"] = int(manifest.get("slide_height") or tokens["slide_height"])
    flags: list[str] = []
    warnings: list[str] = list(plan.get("warnings") or [])
    work = {"chart_rasters": {}}
    prototype_mode = tokens.get("build_mode") == "prototypes" and len(dest.slides) > 0
    original_count = len(dest.slides)

    if not prototype_mode:
        delete_all_slides(dest)

    for entry in plan.get("entries") or []:
        layout_index = int(entry.get("template_layout_index") or layout_index_for_role(tokens, entry.get("role") or "TITLE_ONLY"))
        entry_flags = list(entry.get("flags") or [])
        if prototype_mode:
            if layout_index >= original_count:
                layout_index = 0
                warnings.append(f"Source slide {entry.get('source_index')} fell back to prototype 0.")
            slide = duplicate_slide(dest, layout_index)
            tokens["_layout_index"] = layout_index
            tokens["_current_layout"] = layout_index
            fill_prototype_slide(slide, entry, tokens, brand, entry_flags, work)
        else:
            if layout_index >= len(dest.slide_layouts):
                layout_index = 0
                warnings.append(f"Source slide {entry.get('source_index')} fell back to layout 0.")
            layout = dest.slide_layouts[layout_index]
            slide = dest.slides.add_slide(layout)
            tokens["_layout_index"] = layout_index
            tokens["_current_layout"] = layout_index
            fill_layout_slide(slide, entry, tokens, brand, entry_flags, work)
            cleanup_placeholders(slide)
        remap_srgb_in_slide(slide, tokens.get("srgb_map") or {})
        notes = entry.get("notes") or ""
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        set_slide_hidden(slide, bool(entry.get("hidden")))
        flags.extend(entry_flags)

    if prototype_mode:
        for _ in range(original_count):
            delete_slide(dest, 0)
    patch_theme_part(dest, tokens.get("theme") or {})

    recreate_sections(dest, manifest.get("sections") or [], [])
    output_path = out_dir / "OUTPUT.pptx"
    dest.save(str(output_path))
    unique_flags = sorted(set(flags))
    meta = {
        "output": str(output_path),
        "slide_count": len(plan.get("entries") or []),
        "flags": unique_flags,
        "warnings": warnings,
        "colorway": tokens.get("colorway"),
    }
    dump_json(out_dir / "build_meta.json", meta)
    dump_json(out_dir / "brand_tokens.json", {k: v for k, v in tokens.items() if not str(k).startswith("_")})
    return meta


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build OUTPUT.pptx from plan.json")
    parser.add_argument("--out-dir", default=".")
    parser.add_argument("--source", default=None)
    parser.add_argument("--template", default=None)
    parser.add_argument("--brand-md", default=None)
    parser.add_argument("--colorway", default=None)
    args = parser.parse_args(argv)
    out_dir = Path(args.out_dir)
    tokens = apply_colorway(load_json(out_dir / "brand_tokens.json"), args.colorway)
    dump_json(out_dir / "brand_tokens.json", tokens)
    manifest = load_json(out_dir / "source_manifest.json")
    plan = load_json(out_dir / "plan.json")
    template = Path(args.template) if args.template else Path(tokens.get("template_path") or tokens.get("source_template_path"))
    source = Path(args.source) if args.source else Path(manifest.get("source_path"))
    brand_md = Path(args.brand_md) if args.brand_md else None
    if brand_md is None:
        candidate = source.parent / "brand.md"
        if candidate.exists():
            brand_md = candidate
    meta = build_presentation(
        template_path=template,
        source_path=source,
        out_dir=out_dir,
        tokens=tokens,
        manifest=manifest,
        plan=plan,
        brand_md=brand_md,
        colorway=args.colorway,
    )
    print(f"Wrote {meta['output']} ({meta['slide_count']} slides, colorway={meta.get('colorway')})")
    return 0


if __name__ == "__main__":
    sys.exit(main())

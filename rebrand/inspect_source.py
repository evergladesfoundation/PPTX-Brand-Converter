#!/usr/bin/env python3
"""Inspect a source deck → source_manifest.json."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from helpers import (  # noqa: E402
    blob_hash,
    chart_is_external,
    concatenated_text,
    extract_chart,
    extract_sections,
    extract_smartart_text,
    extract_table,
    has_timing,
    is_edge_logo,
    is_footer_shape,
    is_line_or_arrow,
    iter_shapes_abs,
    paragraphs_from_shape,
    picture_crop,
    placeholder_idx,
    placeholder_type_name,
    shape_kind,
    slide_hidden,
)

TITLE_TYPE_NAMES = {"TITLE", "CENTER_TITLE", "VERTICAL_TITLE"}
SUBTITLE_TYPE_NAMES = {"SUBTITLE"}


def _bounds(shape, left: int, top: int) -> dict[str, int]:
    return {
        "left": int(left),
        "top": int(top),
        "width": int(getattr(shape, "width", 0) or 0),
        "height": int(getattr(shape, "height", 0) or 0),
    }


def _ph_name(shape) -> str | None:
    return placeholder_type_name(shape)


def inspect_shape(shape, left: int, top: int, slide_w: int, slide_h: int, media_dir: Path, slide_index: int, pic_n: int) -> dict[str, Any]:
    kind = shape_kind(shape)
    rec: dict[str, Any] = {
        "id": shape.shape_id,
        "name": shape.name,
        "kind": kind,
        "placeholder_type": _ph_name(shape),
        "placeholder_idx": placeholder_idx(shape),
        "bounds": _bounds(shape, left, top),
        "has_hyperlink": False,
    }
    text_paras = paragraphs_from_shape(shape) if getattr(shape, "has_text_frame", False) else []
    rec["text_frame"] = {
        "paragraph_count": len(text_paras),
        "levels": sorted({p["level"] for p in text_paras}),
        "bullets": any(p.get("bullet") or p.get("numbered") for p in text_paras),
        "paragraphs": text_paras,
    } if text_paras else None
    rec["text"] = concatenated_text(text_paras) if text_paras else ""
    if any(run.get("hyperlink") for p in text_paras for run in p.get("runs") or []):
        rec["has_hyperlink"] = True

    if kind == "picture" or shape.shape_type in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}:
        rec["kind"] = "picture"
        try:
            image = shape.image
            blob = image.blob
            ext = image.ext or "png"
            media_dir.mkdir(parents=True, exist_ok=True)
            path = media_dir / f"slide{slide_index}-img{pic_n}.{ext}"
            path.write_bytes(blob)
            rec["image"] = {
                "path": str(path),
                "ext": ext,
                "bytes": len(blob),
                "hash": blob_hash(blob),
                "content_type": image.content_type,
                "crop": picture_crop(shape),
            }
            try:
                from PIL import Image
                import io

                with Image.open(io.BytesIO(blob)) as im:
                    rec["image"]["px"] = [im.width, im.height]
            except Exception:
                rec["image"]["px"] = [None, None]
        except Exception as exc:
            rec["image"] = {"error": str(exc)}
    if kind == "table":
        rec["table"] = extract_table(shape)
    if kind == "chart":
        rec["chart"] = extract_chart(shape)
        rec["chart"]["external"] = rec["chart"].get("external") or chart_is_external(shape)
    if kind == "other":
        rec["is_line"] = is_line_or_arrow(shape)
    rec["drop_footer"] = is_footer_shape(
        shape,
        rec.get("text") or "",
        slide_w,
        slide_h,
        left,
        top,
        rec["bounds"]["width"],
        rec["bounds"]["height"],
    )
    area = rec["bounds"]["width"] * rec["bounds"]["height"]
    rec["area_ratio"] = area / max(1, slide_w * slide_h)
    if rec["kind"] == "picture":
        rec["edge_logo"] = is_edge_logo(
            area,
            slide_w * slide_h,
            left,
            top,
            rec["bounds"]["width"],
            rec["bounds"]["height"],
            slide_w,
            slide_h,
        )
    return rec


def classify_source_slide(index: int, total: int, title: str, subtitle: str, shapes: list[dict[str, Any]], slide_w: int, slide_h: int) -> str:
    content = [s for s in shapes if not s.get("drop_footer") and s.get("kind") != "group"]
    pictures = [s for s in content if s.get("kind") == "picture"]
    tables = [s for s in content if s.get("kind") == "table"]
    charts = [s for s in content if s.get("kind") == "chart"]
    title_norm = (title or "").strip()
    text_blocks = [
        s
        for s in content
        if s.get("kind") == "text"
        and (s.get("placeholder_type") or "") not in TITLE_TYPE_NAMES | SUBTITLE_TYPE_NAMES
        and (s.get("text") or "").strip()
        and (s.get("text") or "").strip() != title_norm
    ]
    title_words = len((title or "").split())
    body_text = " ".join((s.get("text") or "") for s in text_blocks).strip()
    last_blob = f"{title} {subtitle} {body_text}"

    # Title + subtitle only, or first-slide cover
    if title and subtitle and not body_text and not tables and not charts and not pictures:
        return "TITLE"
    if (
        index == 0
        and not tables
        and not charts
        and not pictures
        and len(text_blocks) <= 1
        and len(body_text) < 280
    ):
        return "TITLE"

    if title_words <= 12 and not body_text and not tables and not charts and not pictures:
        return "SECTION"
    if not title and title_words == 0 and text_blocks:
        only = (text_blocks[0].get("text") or "").strip()
        if len(text_blocks) == 1 and len(only.split()) <= 12 and not tables and not charts and not pictures:
            return "SECTION"

    if index == total - 1 and CONTACT_LIKE(last_blob):
        return "CLOSING"

    if pictures:
        dominating = [p for p in pictures if p.get("area_ratio", 0) > 0.5]
        if dominating and len(text_blocks) <= 2:
            return "PICTURE"

    spatial_blocks = [
        s
        for s in content
        if s.get("kind") in {"text", "picture", "table"}
        and (s.get("placeholder_type") or "") not in TITLE_TYPE_NAMES | SUBTITLE_TYPE_NAMES
        and not s.get("drop_footer")
    ]
    if _two_side_by_side(spatial_blocks, slide_w):
        return "TWO_CONTENT"

    if (tables or charts) and len(text_blocks) <= 1:
        return "TITLE_ONLY"

    if title and len(text_blocks) == 1 and not tables and not charts and not pictures:
        return "TITLE_BODY"
    if len(text_blocks) >= 1 and not tables and not charts and not pictures:
        return "TITLE_BODY"
    return "TITLE_ONLY"


def CONTACT_LIKE(text: str) -> bool:
    from helpers import CONTACT_RE

    return bool(CONTACT_RE.search(text or ""))


def _two_side_by_side(blocks: list[dict[str, Any]], slide_w: int) -> bool:
    if len(blocks) < 2:
        return False
    mid = slide_w / 2
    left = [b for b in blocks if b["bounds"]["left"] + b["bounds"]["width"] / 2 < mid]
    right = [b for b in blocks if b["bounds"]["left"] + b["bounds"]["width"] / 2 >= mid]
    return bool(left) and bool(right) and len(blocks) >= 2


def pick_title(shapes: list[dict[str, Any]]) -> tuple[str, str]:
    title = ""
    subtitle = ""
    for shape in shapes:
        ph = shape.get("placeholder_type") or ""
        text = (shape.get("text") or "").strip()
        if not text:
            continue
        if ph in TITLE_TYPE_NAMES and not title:
            title = text
        elif ph in SUBTITLE_TYPE_NAMES and not subtitle:
            subtitle = text
    if not title:
        named = [s for s in shapes if "title" in (s.get("name") or "").lower() and (s.get("text") or "").strip()]
        if named:
            title = named[0]["text"].strip()
    if not title:
        texts = [s for s in shapes if s.get("kind") == "text" and (s.get("text") or "").strip() and not s.get("drop_footer")]
        texts.sort(key=lambda s: (s["bounds"]["top"], s["bounds"]["left"]))
        if texts:
            candidate = texts[0]["text"].strip()
            if len(candidate.split()) <= 16:
                title = candidate
    return title, subtitle


def inspect_source(source_path: Path, out_dir: Path) -> dict[str, Any]:
    out_dir.mkdir(parents=True, exist_ok=True)
    media_dir = out_dir / "source_media"
    prs = Presentation(str(source_path))
    slide_w = int(prs.slide_width)
    slide_h = int(prs.slide_height)
    smartart_pool = extract_smartart_text(source_path)
    slides = []
    picture_hashes: dict[str, int] = {}
    for index, slide in enumerate(prs.slides):
        layout_name = slide.slide_layout.name if slide.slide_layout is not None else ""
        pic_n = 0
        shapes = []
        for shape, left, top in iter_shapes_abs(slide.shapes):
            rec = inspect_shape(shape, left, top, slide_w, slide_h, media_dir, index, pic_n)
            if rec.get("kind") == "picture":
                pic_n += 1
                digest = (rec.get("image") or {}).get("hash")
                if digest:
                    picture_hashes[digest] = picture_hashes.get(digest, 0) + 1
            if rec.get("kind") == "smartart" and not rec.get("text"):
                rec["smartart_text"] = smartart_pool
            shapes.append(rec)
        title, subtitle = pick_title(shapes)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text or ""
        role = classify_source_slide(index, len(prs.slides), title, subtitle, shapes, slide_w, slide_h)
        slides.append(
            {
                "index": index,
                "layout_name": layout_name,
                "title_text": title,
                "subtitle_text": subtitle,
                "role": role,
                "shapes": shapes,
                "notes_text": notes,
                "has_transitions_or_animations": has_timing(slide),
                "hidden": slide_hidden(slide),
            }
        )

    slide_count = max(1, len(slides))
    repeated = {h for h, n in picture_hashes.items() if n / slide_count >= 0.60}
    for slide in slides:
        for shape in slide["shapes"]:
            digest = (shape.get("image") or {}).get("hash")
            shape["repeated_logo"] = bool(digest and digest in repeated)
            if shape.get("kind") == "picture" and (shape.get("edge_logo") or shape.get("repeated_logo")):
                shape["drop_logo"] = True
            else:
                shape["drop_logo"] = False

    manifest = {
        "source_path": str(source_path.resolve()),
        "slide_width": slide_w,
        "slide_height": slide_h,
        "slide_count": len(slides),
        "sections": extract_sections(prs),
        "repeated_logo_hashes": sorted(repeated),
        "slides": slides,
    }
    from helpers import dump_json

    dump_json(out_dir / "source_manifest.json", manifest)
    return manifest


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Inspect a source PPTX into source_manifest.json")
    parser.add_argument("source")
    parser.add_argument("--out-dir", default=".")
    args = parser.parse_args(argv)
    source = Path(args.source)
    if not source.exists():
        print(f"Source not found: {source}", file=sys.stderr)
        return 1
    manifest = inspect_source(source, Path(args.out_dir))
    print(f"Wrote {Path(args.out_dir) / 'source_manifest.json'} ({manifest['slide_count']} slides)")
    return 0


if __name__ == "__main__":
    sys.exit(main())

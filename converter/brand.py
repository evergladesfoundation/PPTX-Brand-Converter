"""Everglades Foundation colors and chrome helpers for generated decks."""

from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

MARSH = RGBColor(0x0B, 0x43, 0x3F)
LAGOON = RGBColor(0x08, 0x93, 0xA6)
SAWGRASS = RGBColor(0xC1, 0xEA, 0x40)
INK = RGBColor(0x1C, 0x1C, 0x1C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FOOTER_TEXT = "The Everglades Foundation"


def _no_line(shape) -> None:
    shape.line.fill.background()


def add_brand_chrome(slide, prs, *, dark: bool = False) -> None:
    """Sawgrass top bar, marsh footer, Foundation wordmark."""
    width = prs.slide_width
    height = prs.slide_height
    footer_h = Inches(0.38)
    bar_h = Inches(0.1)

    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, width, bar_h)
    top.fill.solid()
    top.fill.fore_color.rgb = SAWGRASS
    _no_line(top)

    footer = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, height - footer_h, width, footer_h
    )
    footer.fill.solid()
    footer.fill.fore_color.rgb = MARSH
    _no_line(footer)

    label = slide.shapes.add_textbox(
        Inches(0.4),
        height - footer_h,
        width - Inches(0.8),
        footer_h,
    )
    tf = label.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = FOOTER_TEXT
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.LEFT

    if dark:
        fill = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, bar_h, width, height - footer_h - bar_h)
        fill.fill.solid()
        fill.fill.fore_color.rgb = MARSH
        _no_line(fill)
        sp_tree = slide.shapes._spTree
        element = fill._element
        sp_tree.remove(element)
        sp_tree.insert(2, element)

"""Deprecated CLI. Conversion lives in rebrand/ (inspect → plan → build → QA).

Kept so existing scripts do not break; the staff app no longer calls this file.
Do not extend this module to satisfy the rebrand spec.
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from brand import LAGOON, MARSH, WHITE, add_brand_chrome  # noqa: E402
from charts import (  # noqa: E402
    chart_is_external,
    chart_shapes,
    copy_chart,
    default_chart_box,
    iter_shapes,
)
from layouts import LAYOUT_IDS, LABELS, PPTX_LAYOUT_NAME  # noqa: E402

TEMPLATE_PATH = ROOT.parent / "templates" / "everglades.pptx"
TITLE_TYPES = {
    PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.CENTER_TITLE,
    PP_PLACEHOLDER.VERTICAL_TITLE,
}
BODY_TYPES = {
    PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.OBJECT,
    PP_PLACEHOLDER.VERTICAL_BODY,
    PP_PLACEHOLDER.SUBTITLE,
}
PICTURE_TYPES = {PP_PLACEHOLDER.PICTURE, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.MEDIA_CLIP}


def layout_by_name(prs: Presentation, name: str):
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    raise KeyError(f"Layout not found: {name}")


def delete_all_slides(prs: Presentation) -> None:
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)


def _placeholder_type(shape) -> Any | None:
    try:
        return shape.placeholder_format.type
    except Exception:
        return None


def extract_slide(slide, index: int, work_dir: Path) -> dict[str, Any]:
    title = ""
    bodies: list[str] = []
    pictures: list[str] = []
    tables: list[list[list[str]]] = []
    warnings: list[str] = []
    notes = ""
    has_chart = False
    external_chart = False

    for shape in iter_shapes(slide.shapes):
        ph_type = _placeholder_type(shape)
        if getattr(shape, "has_chart", False) or shape.shape_type == MSO_SHAPE_TYPE.CHART:
            has_chart = True
            if chart_is_external(shape):
                external_chart = True
                warnings.append("Chart is linked to an external Excel file and cannot stay live.")
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            pictures.append(_save_picture(shape, work_dir, index, len(pictures)))
            continue
        if getattr(shape, "has_table", False):
            table = shape.table
            tables.append(
                [[cell.text.strip() for cell in row.cells] for row in table.rows]
            )
            continue
        if not getattr(shape, "has_text_frame", False):
            continue
        text = "\n".join(
            paragraph.text.strip()
            for paragraph in shape.text_frame.paragraphs
            if paragraph.text.strip()
        ).strip()
        if not text:
            continue
        if ph_type in TITLE_TYPES or (not title and "title" in (shape.name or "").lower()):
            if not title:
                title = text
            else:
                bodies.append(text)
        else:
            bodies.append(text)

    if slide.has_notes_slide:
        notes = (slide.notes_slide.notes_text_frame.text or "").strip()

    body_lines = _split_body(bodies)
    suggested = classify(index, title, body_lines, pictures, has_chart)
    return {
        "index": index,
        "title": title,
        "body": body_lines,
        "bodyPreview": "\n".join(body_lines[:8]),
        "hasImage": bool(pictures),
        "imageCount": len(pictures),
        "hasChart": has_chart,
        "externalChart": external_chart,
        "hasTable": bool(tables),
        "notes": notes,
        "suggestedLayout": suggested,
        "warnings": warnings,
        "_pictures": pictures,
        "_tables": tables,
    }


def _save_picture(shape, work_dir: Path, slide_index: int, pic_index: int) -> str:
    image = shape.image
    ext = image.ext or "png"
    path = work_dir / f"slide{slide_index}-img{pic_index}.{ext}"
    path.write_bytes(image.blob)
    return str(path)


def _split_body(bodies: list[str]) -> list[str]:
    lines: list[str] = []
    for block in bodies:
        for line in block.replace("\r", "").split("\n"):
            cleaned = line.strip()
            if cleaned:
                lines.append(cleaned)
    return lines


def classify(
    index: int,
    title: str,
    body: list[str],
    pictures: list[str],
    has_chart: bool,
) -> str:
    if has_chart:
        return "chart"
    joined = " ".join(body)
    if index == 0 and len(body) <= 3 and len(joined) < 280:
        return "title"
    if pictures and len(body) <= 2:
        return "image"
    headline = (title or (body[0] if body else "")).strip()
    if index > 0 and not pictures and len(body) <= 1 and len(headline) < 48:
        return "section"
    stripped = (title or joined).strip()
    if stripped.startswith(("\"", "“", "‘", "'")):
        return "quote"
    if _looks_two_column(body):
        return "two_column"
    return "content"


def _looks_two_column(body: list[str]) -> bool:
    if len(body) < 4:
        return False
    mid = len(body) // 2
    return abs(len(" ".join(body[:mid])) - len(" ".join(body[mid:]))) < 80


def parse_presentation(path: Path) -> dict[str, Any]:
    work_dir = Path(tempfile.mkdtemp(prefix="ef-parse-"))
    try:
        prs = Presentation(path)
        slides = [extract_slide(slide, i, work_dir) for i, slide in enumerate(prs.slides)]
        public = []
        for slide in slides:
            item = {k: v for k, v in slide.items() if not k.startswith("_")}
            public.append(item)
        return {
            "fileName": path.name,
            "slideCount": len(public),
            "layouts": [{"id": key, "label": LABELS[key]} for key in LAYOUT_IDS],
            "slides": public,
        }
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def _set_text(shape, text: str, *, color=None, size: int | None = None, bold: bool | None = None) -> None:
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run.font.name = "Poppins"


def _set_bullets(shape, lines: list[str]) -> None:
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.clear()
    if not lines:
        return
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.font.name = "Poppins"
        p.font.size = Pt(16)
        p.font.color.rgb = MARSH


def _find_placeholder(slide, types: set) -> Any | None:
    for shape in slide.placeholders:
        try:
            if shape.placeholder_format.type in types:
                return shape
        except Exception:
            continue
    return None


def _find_placeholders(slide, types: set) -> list[Any]:
    found = []
    for shape in slide.placeholders:
        try:
            if shape.placeholder_format.type in types:
                found.append(shape)
        except Exception:
            continue
    return found


def _title_shape(slide):
    return _find_placeholder(slide, TITLE_TYPES) or _find_placeholder(
        slide, {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
    )


def _add_table(slide, prs, data: list[list[str]]) -> None:
    if not data:
        return
    rows = len(data)
    cols = max(len(row) for row in data)
    table_shape = slide.shapes.add_table(
        rows,
        cols,
        Inches(0.7),
        Inches(4.6),
        prs.slide_width - Inches(1.4),
        Inches(1.8),
    )
    table = table_shape.table
    for r, row in enumerate(data):
        for c in range(cols):
            table.cell(r, c).text = row[c] if c < len(row) else ""


def fill_slide(
    dest,
    prs: Presentation,
    source_slide,
    extracted: dict[str, Any],
    layout_id: str,
) -> list[str]:
    warnings: list[str] = list(extracted.get("warnings") or [])
    title = extracted.get("title") or ""
    body: list[str] = list(extracted.get("body") or [])
    pictures: list[str] = list(extracted.get("_pictures") or [])
    tables = list(extracted.get("_tables") or [])
    dark = layout_id == "title"
    add_brand_chrome(dest, prs, dark=dark)

    title_shape = _title_shape(dest)
    if title_shape:
        _set_text(
            title_shape,
            title or (body[0] if body else ""),
            color=WHITE if dark else MARSH,
            size=36 if layout_id in {"title", "quote"} else 28,
            bold=True,
        )

    if layout_id == "two_column":
        mids = max(1, len(body) // 2)
        columns = [body[:mids], body[mids:]]
        bodies = _find_placeholders(dest, BODY_TYPES)
        for i, box in enumerate(bodies[:2]):
            _set_bullets(box, columns[i] if i < len(columns) else [])
    elif layout_id == "image":
        pic = _find_placeholder(dest, PICTURE_TYPES)
        if pic and pictures:
            try:
                pic.insert_picture(pictures[0])
            except Exception:
                dest.shapes.add_picture(
                    pictures[0],
                    Inches(0.7),
                    Inches(1.6),
                    width=prs.slide_width - Inches(1.4),
                )
        caption = _find_placeholder(dest, BODY_TYPES)
        if caption:
            _set_text(caption, "\n".join(body[:3]), color=LAGOON, size=14, bold=False)
    elif layout_id == "quote":
        quote = title if title else " ".join(body)
        if title_shape:
            _set_text(title_shape, quote, color=MARSH, size=28, bold=True)
    elif layout_id == "chart":
        pass
    else:
        body_shape = _find_placeholder(dest, BODY_TYPES)
        if body_shape:
            if layout_id == "title":
                _set_text(body_shape, "\n".join(body[:4]), color=WHITE, size=18, bold=False)
            elif layout_id == "section":
                _set_text(body_shape, "\n".join(body[:2]), color=LAGOON, size=18, bold=False)
            else:
                _set_bullets(body_shape, body)

    if layout_id == "chart":
        left, top, width, height = default_chart_box(prs)
        copied = False
        for shape in chart_shapes(source_slide):
            ok = copy_chart(shape, dest, left=left, top=top, width=width, height=height)
            copied = copied or ok
            if not ok:
                warnings.append("A chart could not be copied as a live Excel object.")
        if not copied and pictures:
            dest.shapes.add_picture(pictures[0], left, top, width=width)
            warnings.append("No live chart was copied; used an image instead.")
        elif not copied:
            warnings.append("No live chart was found on this slide.")

    if layout_id == "content" and pictures and not chart_shapes(source_slide):
        dest.shapes.add_picture(
            pictures[0],
            Inches(8.6),
            Inches(1.6),
            width=Inches(4.0),
        )

    if tables and layout_id == "content":
        _add_table(dest, prs, tables[0])

    notes = extracted.get("notes") or ""
    if notes:
        dest.notes_slide.notes_text_frame.text = notes

    return warnings


def convert_presentation(source_path: Path, dest_path: Path, layout_ids: list[str]) -> dict[str, Any]:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"Missing {TEMPLATE_PATH}. Run converter/build_template.py first."
        )
    work_dir = Path(tempfile.mkdtemp(prefix="ef-convert-"))
    try:
        source = Presentation(source_path)
        extracted = [
            extract_slide(slide, i, work_dir) for i, slide in enumerate(source.slides)
        ]
        if len(layout_ids) != len(extracted):
            raise ValueError(
                f"Expected {len(extracted)} layout choices, received {len(layout_ids)}."
            )
        for layout_id in layout_ids:
            if layout_id not in LAYOUT_IDS:
                raise ValueError(f"Unknown layout: {layout_id}")

        dest = Presentation(str(TEMPLATE_PATH))
        delete_all_slides(dest)
        warnings: list[str] = []
        for i, slide_data in enumerate(extracted):
            layout_id = layout_ids[i]
            layout = layout_by_name(dest, PPTX_LAYOUT_NAME[layout_id])
            new_slide = dest.slides.add_slide(layout)
            slide_warnings = fill_slide(
                new_slide, dest, source.slides[i], slide_data, layout_id
            )
            warnings.extend(f"Slide {i + 1}: {msg}" for msg in slide_warnings)

        dest_path.parent.mkdir(parents=True, exist_ok=True)
        dest.save(dest_path)
        return {
            "output": str(dest_path),
            "slideCount": len(extracted),
            "warnings": warnings,
        }
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def _cli(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Everglades PPTX converter")
    sub = parser.add_subparsers(dest="cmd", required=True)

    parse_cmd = sub.add_parser("parse", help="Extract slide content as JSON")
    parse_cmd.add_argument("input")

    convert_cmd = sub.add_parser("convert", help="Rebuild onto the Everglades template")
    convert_cmd.add_argument("input")
    convert_cmd.add_argument("output")
    convert_cmd.add_argument(
        "--layouts",
        required=True,
        help="JSON array of layout ids, or path to a JSON file",
    )

    args = parser.parse_args(argv)
    try:
        if args.cmd == "parse":
            result = parse_presentation(Path(args.input))
            json.dump(result, sys.stdout, ensure_ascii=False)
            sys.stdout.write("\n")
            return 0
        layouts_raw = args.layouts
        layouts_path = Path(layouts_raw)
        if layouts_path.exists():
            layout_ids = json.loads(layouts_path.read_text(encoding="utf-8"))
        else:
            layout_ids = json.loads(layouts_raw)
        result = convert_presentation(Path(args.input), Path(args.output), layout_ids)
        json.dump(result, sys.stdout, ensure_ascii=False)
        sys.stdout.write("\n")
        return 0
    except Exception as exc:
        json.dump({"error": str(exc)}, sys.stderr)
        sys.stderr.write("\n")
        return 1


if __name__ == "__main__":
    sys.exit(_cli())

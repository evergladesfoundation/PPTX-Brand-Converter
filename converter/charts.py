"""Copy live Office charts (chart XML + embedded Excel) into a destination slide."""

from __future__ import annotations

from copy import deepcopy
from typing import Any

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.parts.chart import ChartPart
from pptx.parts.embeddedpackage import EmbeddedXlsxPart
from pptx.util import Inches

C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def iter_shapes(shapes) -> list[Any]:
    found = []
    for shape in shapes:
        found.append(shape)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            found.extend(iter_shapes(shape.shapes))
    return found


def chart_shapes(slide) -> list[Any]:
    charts = []
    for shape in iter_shapes(slide.shapes):
        has_chart = getattr(shape, "has_chart", False)
        if has_chart or shape.shape_type == MSO_SHAPE_TYPE.CHART:
            charts.append(shape)
    return charts


def chart_is_external(shape) -> bool:
    try:
        xlsx = shape.chart.part.chart_workbook.xlsx_part
    except Exception:
        return True
    return xlsx is None


def _copy_related_xlsx(source_chart_part: ChartPart, dest_chart_part: ChartPart) -> None:
    try:
        xlsx_part = source_chart_part.chart_workbook.xlsx_part
    except Exception:
        return
    if xlsx_part is None:
        return
    dest_chart_part.chart_workbook.xlsx_part = EmbeddedXlsxPart.new(
        xlsx_part.blob,
        dest_chart_part.package,
    )


def copy_chart(source_shape, dest_slide, *, left, top, width, height) -> bool:
    """Clone a chart graphicFrame and its embedded workbook into dest_slide.

    Returns False when the chart is linked to an external Excel file.
    """
    if chart_is_external(source_shape):
        return False

    source_part: ChartPart = source_shape.chart.part
    dest_slide_part = dest_slide.part
    package = dest_slide_part.package

    partname = package.next_partname(ChartPart.partname_template)
    new_chart_part = ChartPart(
        partname,
        source_part.content_type,
        package,
        deepcopy(source_part._element),
    )
    _copy_related_xlsx(source_part, new_chart_part)
    r_id = dest_slide_part.relate_to(new_chart_part, RT.CHART)

    graphic = deepcopy(source_shape._element)
    chart_ref = graphic.find(f".//{{{C_NS}}}chart")
    if chart_ref is None:
        return False
    chart_ref.set(f"{{{R_NS}}}id", r_id)

    _set_xfrm(graphic, left, top, width, height)
    dest_slide.shapes._spTree.append(graphic)
    return True


def _set_xfrm(graphic, left, top, width, height) -> None:
    xfrm = graphic.find(f".//{{{A_NS}}}xfrm")
    if xfrm is None:
        xfrm = graphic.find(f".//{{{P_NS}}}xfrm")
    if xfrm is None:
        return
    off = xfrm.find(f"{{{A_NS}}}off")
    ext = xfrm.find(f"{{{A_NS}}}ext")
    if off is not None:
        off.set("x", str(int(left)))
        off.set("y", str(int(top)))
    if ext is not None:
        ext.set("cx", str(int(width)))
        ext.set("cy", str(int(height)))


def default_chart_box(prs):
    left = Inches(0.7)
    top = Inches(1.55)
    width = prs.slide_width - Inches(1.4)
    height = prs.slide_height - Inches(2.15)
    return left, top, width, height

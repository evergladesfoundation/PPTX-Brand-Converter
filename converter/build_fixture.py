"""Build fixtures/sample.pptx — mixed layouts plus a live clustered-column chart."""

from __future__ import annotations

import struct
import sys
import zlib
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from brand import LAGOON, MARSH

ROOT = Path(__file__).resolve().parent.parent
FIXTURE = ROOT / "fixtures" / "sample.pptx"


def write_png(path: Path, rgb: tuple[int, int, int] = (8, 147, 166)) -> None:
    width = height = 120
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    png = (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw, 9))
        + chunk(b"IEND", b"")
    )
    path.write_bytes(png)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]

    title = prs.slides.add_slide(title_layout)
    title.shapes.title.text = "Off-brand staff briefing"
    subtitle = title.placeholders[1]
    subtitle.text = "Calibri, default theme, mixed content"

    section = prs.slides.add_slide(blank)
    box = section.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(11), Inches(1.5))
    p = box.text_frame.paragraphs[0]
    p.text = "PROGRAM UPDATE"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = MARSH

    content = prs.slides.add_slide(content_layout)
    content.shapes.title.text = "Priorities this quarter"
    body = content.placeholders[1].text_frame
    body.paragraphs[0].text = "Complete CERP project reviews"
    line = body.add_paragraph()
    line.text = "Brief county partners on water quality"
    line = body.add_paragraph()
    line.text = "Publish the restoration scorecard"

    image_slide = prs.slides.add_slide(blank)
    png = ROOT / "fixtures" / "swatch.png"
    png.parent.mkdir(parents=True, exist_ok=True)
    write_png(png, (11, 67, 63))
    image_slide.shapes.add_picture(str(png), Inches(3.2), Inches(1.4), width=Inches(6.8))
    cap = image_slide.shapes.add_textbox(Inches(0.8), Inches(6.2), Inches(11.5), Inches(0.6))
    cap.text_frame.paragraphs[0].text = "Marsh survey site"

    chart_slide = prs.slides.add_slide(blank)
    heading = chart_slide.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(12), Inches(0.8))
    hp = heading.text_frame.paragraphs[0]
    hp.text = "Restoration investment"
    hp.font.size = Pt(28)
    hp.font.bold = True
    hp.font.color.rgb = LAGOON
    data = CategoryChartData()
    data.categories = ["2022", "2023", "2024", "2025"]
    data.add_series("Funding ($M)", (18.4, 21.0, 24.5, 27.2))
    chart_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1.2),
        Inches(1.5),
        Inches(10.8),
        Inches(5.0),
        data,
    )

    quote = prs.slides.add_slide(blank)
    q = quote.shapes.add_textbox(Inches(1.2), Inches(2.6), Inches(10.8), Inches(2.2))
    qp = q.text_frame.paragraphs[0]
    qp.text = '"A thriving Everglades means a thriving economy."'
    qp.font.size = Pt(28)
    qp.font.italic = True
    qp.font.color.rgb = MARSH

    notes_slide = content.notes_slide
    notes_slide.notes_text_frame.text = "Speaker note: keep the scorecard live."

    FIXTURE.parent.mkdir(parents=True, exist_ok=True)
    prs.save(FIXTURE)
    return FIXTURE


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
    sys.exit(0)
